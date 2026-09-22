"""Create a compact, reproducible workbook for qualitative error analysis."""

import argparse
import hashlib
import json
import random
import re
from collections import Counter
from pathlib import Path

from meeting_qa_chunking.config import load_run_config
from meeting_qa_chunking.evidence import render_gold_evidence
from meeting_qa_chunking.evidence_preparation import prepare_retrieved_evidence
from meeting_qa_chunking.judging import parse_axis_judgment
from meeting_qa_chunking.qmsum import load_meeting


QUERY_TYPES = ("synthesis", "participant-specific", "reasoning/response")
ANNOTATION_CHUNKERS = ("turn_packed", "lumber")
ANNOTATION_BUDGET = 2048
MAX_EXCEL_TEXT = 32_000


def read_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def index_records(records, *fields):
    indexed = {tuple(row[field] for field in fields): row for row in records}
    if len(indexed) != len(records):
        raise ValueError(f"Duplicate records indexed by {fields}")
    return indexed


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def speaker_aliases(speakers: set[str]) -> set[str]:
    aliases = {speaker.casefold() for speaker in speakers if len(speaker) > 2}
    for speaker in speakers:
        aliases.update(
            item.casefold()
            for item in re.findall(r"\(([^)]+)\)", speaker)
            if len(item) > 3
        )
    return aliases


def classify_query(question: str, speakers: set[str]) -> str:
    """Assign a transparent lexical query type used only for sampling."""

    text = question.casefold()
    if any(alias in text for alias in speaker_aliases(speakers)):
        return "participant-specific"
    reasoning_markers = (
        "why ", "how ", "reason", "respond", "response", "reaction",
        "agree", "disagree", "decide", "decision", "problem", "solution",
        "resolve", "outcome", "cause", "effect", "difference", "compare",
        "relationship", "purpose", "implication", "challenge", "concern",
    )
    if any(marker in text for marker in reasoning_markers):
        return "reasoning/response"
    return "synthesis"


def condition_names(retriever: str) -> list[str]:
    return [
        f"{chunker}__{retriever}__w{ANNOTATION_BUDGET}"
        for chunker in ANNOTATION_CHUNKERS
    ]


def load_cases(run, answer_stage: str, retriever: str) -> list[dict]:
    conditions = condition_names(retriever)
    evaluation = index_records(
        read_json(run.evaluation_dir / f"{answer_stage}.json")["records"],
        "meeting_id",
        "question_index",
        "condition",
    )
    cases = []
    for meeting_id in run.meeting_ids():
        meeting = load_meeting(run.data_dir / f"{meeting_id}.json")
        retrieval = read_json(run.retrieval_dir / f"{meeting_id}.json")
        answers = read_json(
            run.answers_dir / answer_stage / f"{meeting_id}.json"
        )["questions"]
        oracle_answers = read_json(
            run.answers_dir / "oracle-14b" / f"{meeting_id}.json"
        )["questions"]
        missing = set(conditions) - set(retrieval["configurations"])
        if missing:
            raise ValueError(f"Missing conditions in {meeting_id}: {sorted(missing)}")
        if not (
            len(retrieval["questions"])
            == len(answers)
            == len(oracle_answers)
            == len(meeting.questions)
        ):
            raise ValueError(f"Question-count mismatch in {meeting_id}")

        speakers = {turn.speaker for turn in meeting.turns}
        for index, question in enumerate(meeting.questions):
            if oracle_answers[index]["question_index"] != index:
                raise ValueError(f"Oracle question mismatch in {meeting_id}")
            results = {}
            for name in conditions:
                retrieved = retrieval["questions"][index]["results"][name]
                generated = answers[index]["results"][name]
                judged = evaluation[(meeting_id, index, name)]
                results[name] = {
                    **retrieval["configurations"][name],
                    "precision": retrieved["precision"],
                    "recall": retrieved["recall"],
                    "rougeL": generated["rouge_f1"]["rougeL"],
                    "bertscore_f1": judged["bertscore"]["f1"],
                    "judge": judged["judge"]["score"],
                    "judge_reason": judged["judge"]["reason"],
                    "answer": generated["answer"],
                }
            cases.append(
                {
                    "meeting_id": meeting_id,
                    "question_index": index,
                    "question": question.text,
                    "reference_answer": question.reference_answer,
                    "oracle_answer": oracle_answers[index]["results"]["oracle"]["answer"],
                    "query_type": classify_query(question.text, speakers),
                    "gold_turn_ranges": question.relevant_turn_ranges,
                    "results": results,
                }
            )
    return cases


def add_evidence_aware_judgments(selected: list[dict], path: Path, retriever: str) -> None:
    """Join the saved paired judgments without changing the question sample."""

    records = index_records(
        read_json(path)["records"], "meeting_id", "question_index", "condition"
    )
    for case in selected:
        for condition in condition_names(retriever):
            result = case["results"][condition]
            key = (case["meeting_id"], case["question_index"], condition)
            record = records[key]
            if (
                record["candidate_answer"] != result["answer"]
                or record["retrieved_evidence"] != result["retrieved_evidence"]
                or record["reference_answer"] != case["reference_answer"]
            ):
                raise ValueError(f"Judgment inputs differ from workbook: {key}")
            aware = record["evidence_aware_judge"]
            result["faithfulness"] = parse_axis_judgment(
                aware["faithfulness"]["raw_response"], "unsupported_claims"
            )
            result["sufficiency"] = parse_axis_judgment(
                aware["evidence_sufficiency"]["raw_response"],
                "unavailable_reference_points",
            )


def add_sampling_deltas(cases: list[dict], retriever: str) -> None:
    """Add only the paired effects used by the explicit sampling strata."""

    for case in cases:
        lumber = case["results"][
            f"lumber__{retriever}__w{ANNOTATION_BUDGET}"
        ]
        baseline = case["results"][
            f"turn_packed__{retriever}__w{ANNOTATION_BUDGET}"
        ]
        case["recall_delta"] = lumber["recall"] - baseline["recall"]
        case["judge_delta"] = lumber["judge"] - baseline["judge"]


def _take(
    candidates: list[dict],
    count: int,
    selected_keys: set[tuple[str, int]],
    meeting_counts: Counter,
    reason: str,
    max_per_meeting: int,
) -> list[dict]:
    chosen = []
    for enforce_cap in (True, False):
        for case in candidates:
            key = (case["meeting_id"], case["question_index"])
            if key in selected_keys:
                continue
            local_count = sum(
                item["meeting_id"] == case["meeting_id"] for item in chosen
            )
            if (
                enforce_cap
                and meeting_counts[case["meeting_id"]] + local_count
                >= max_per_meeting
            ):
                continue
            chosen.append(case)
            selected_keys.add(key)
            if len(chosen) == count:
                break
        if len(chosen) == count:
            break
    if len(chosen) != count:
        raise ValueError(f"Could not select {count} {reason} cases")
    for case in chosen:
        meeting_counts[case["meeting_id"]] += 1
        case["selection_reason"] = reason
    return chosen


def select_cases(
    cases: list[dict],
    per_type: int,
    seed: int,
    max_per_meeting: int = 2,
) -> list[dict]:
    """Select explicit paired contrasts plus a seeded random reference set."""

    if per_type < 6:
        raise ValueError("per_type must be at least 6")
    stratum_count = max(1, round(per_type * 0.2))
    random_count = per_type - 3 * stratum_count
    rng = random.Random(seed)
    tie_break = {
        (case["meeting_id"], case["question_index"]): rng.random()
        for case in cases
    }
    selected = []
    selected_keys = set()
    meeting_counts = Counter()

    for query_type in QUERY_TYPES:
        pool = [case for case in cases if case["query_type"] == query_type]
        if len(pool) < per_type:
            raise ValueError(f"Only {len(pool)} {query_type} questions available")
        lumber_wins = sorted(
            (case for case in pool if case["recall_delta"] > 0),
            key=lambda case: (
                -case["recall_delta"],
                tie_break[(case["meeting_id"], case["question_index"])],
            ),
        )
        chosen = _take(
            lumber_wins,
            stratum_count,
            selected_keys,
            meeting_counts,
            "largest Lumber recall advantage",
            max_per_meeting,
        )
        baseline_wins = sorted(
            (case for case in pool if case["recall_delta"] < 0),
            key=lambda case: (
                case["recall_delta"],
                tie_break[(case["meeting_id"], case["question_index"])],
            ),
        )
        chosen += _take(
            baseline_wins,
            stratum_count,
            selected_keys,
            meeting_counts,
            "largest turn-packed recall advantage",
            max_per_meeting,
        )
        non_proportional = sorted(
            (
                case
                for case in pool
                if case["recall_delta"] > 0 and case["judge_delta"] <= 0
            ),
            key=lambda case: (
                -case["recall_delta"],
                case["judge_delta"],
                tie_break[(case["meeting_id"], case["question_index"])],
            ),
        )
        chosen += _take(
            non_proportional,
            stratum_count,
            selected_keys,
            meeting_counts,
            "Lumber recall gain without judge gain",
            max_per_meeting,
        )
        remaining = [
            case
            for case in pool
            if (case["meeting_id"], case["question_index"]) not in selected_keys
        ]
        rng.shuffle(remaining)
        chosen += _take(
            remaining,
            random_count,
            selected_keys,
            meeting_counts,
            "seeded random",
            max_per_meeting,
        )
        selected.extend(chosen)

    for index, case in enumerate(selected, 1):
        case["sample_id"] = f"Q{index:02d}"
    return selected


def hydrate_evidence(run, selected: list[dict], conditions: list[str]) -> None:
    by_meeting = {}
    for case in selected:
        by_meeting.setdefault(case["meeting_id"], []).append(case)
    for meeting_id, meeting_cases in by_meeting.items():
        meeting = load_meeting(run.data_dir / f"{meeting_id}.json")
        _, evidence = prepare_retrieved_evidence(
            meeting,
            run.retrieval_dir / f"{meeting_id}.json",
            run.lumber_dir,
            conditions,
        )
        for case in meeting_cases:
            question = meeting.questions[case["question_index"]]
            case["gold_evidence"], _ = render_gold_evidence(question, meeting)
            for condition in conditions:
                case["results"][condition]["retrieved_evidence"] = evidence[
                    case["question_index"]
                ][condition]["text"]


def excel_text(value) -> str:
    text = "" if value is None else str(value)
    if len(text) <= MAX_EXCEL_TEXT:
        return text
    return text[: MAX_EXCEL_TEXT - 80] + "\n[TRUNCATED AT EXCEL CELL LIMIT]"


def add_table(sheet, headers, rows, widths, name):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [excel_text(value) if isinstance(value, str) else value for value in row]
        )
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[1].height = 32
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(
        displayName=name,
        ref=f"A1:{sheet.cell(1, len(headers)).column_letter}{len(rows) + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def write_instructions(workbook, payload):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet("Instructions")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 115
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Manual qualitative review instructions"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    rows = (
        ("Purpose", "Explain why retrieval differences do or do not translate into answer-quality differences. This is a purposive diagnostic analysis, not an estimate of error prevalence."),
        ("Comparison", "Each sampled question has two adjacent rows: turn-packed and Lumber, both using dense retrieval and a 2,048-word evidence budget."),
        ("Questions sheet", "Contains the 30 sampled questions, query type, reference and oracle-14B answers, annotated gold evidence, and sampling reason."),
        ("Annotations sheet", "Contains 60 paired rows. The independent automatic faithfulness and sufficiency judgments are appended after the six manual fields so existing annotation columns stay fixed."),
        ("Summary sheet", "Updates from manual annotations when opened in Excel and also compares the saved automatic faithfulness and sufficiency scores."),
        ("Review step 1", "Read the question, reference answer, and gold evidence on Questions. Gold evidence is QMSum's annotated relevant turns and may itself be imperfect."),
        ("Review step 2", "For each paired question, inspect retrieved evidence before judging the answer. Optionally hide automatic metric columns I:O and Y:AD during the first pass to reduce anchoring."),
        ("Review step 3", "Complete all six manual fields. Use Notes for the concrete missing, distracting, unsupported, or incorrectly generated detail."),
        ("Evidence sufficient", "Yes = enough evidence for a complete answer; Partial = useful but missing important information; No = insufficient."),
        ("Noise", "Low = little irrelevant material; Med = noticeable but still usable; High = distracting or potentially misleading."),
        ("Answer quality", "1 = invalid or incorrect; 2 = partially correct or incomplete; 3 = correct and sufficiently complete."),
        ("Unsupported content", "Yes only if the generated answer makes a substantive claim not supported by its retrieved evidence."),
        ("Oracle answer", "The Qwen2.5-14B answer generated from gold evidence; compare it with retrieved-evidence answers, but do not treat it as ground truth."),
        ("Gold-only judge", "The original 1-3 correctness/completeness score uses reference answer and gold evidence, not retrieved evidence."),
        ("Automatic faithfulness", "Independent 1-3 judge of whether the generated answer's claims are supported by the retrieved evidence; listed unsupported claims are diagnostic, not manual labels."),
        ("Automatic sufficiency", "Independent 1-3 judge of whether retrieved evidence contains the important reference-answer points, regardless of the generated answer."),
        ("Manual versus automatic", "Complete the six manual fields independently before comparing them with the automatic judgments appended to the right."),
        ("Main error: Retrieval", "Relevant information is absent or too incomplete to answer."),
        ("Main error: Noise", "Relevant evidence is present, but irrelevant context appears to distract or mislead the model."),
        ("Main error: Generation", "Evidence is adequate, but the answer omits, distorts, attributes incorrectly, or invents information."),
        ("Main error: Annotation", "The reference answer or gold evidence appears incomplete, misleading, or incorrect."),
        ("Main error: None", "No material error is apparent."),
        ("Synthesis", "Asks for a topic-level account spanning a discussion or meeting segment."),
        ("Participant-specific", "Names a meeting speaker and asks about that speaker's contribution or view."),
        ("Reasoning/response", "Asks why/how, or about a decision, reaction, problem, solution, or outcome."),
        ("Sampling", f"Ten questions per query type: two largest Lumber recall advantages, two largest turn-packed recall advantages, two Lumber recall gains without a judge-score gain, and four seeded-random questions from the remainder. All comparisons use dense retrieval and 2,048 words. Seed {payload['seed']}."),
        ("No composite score", "Sampling uses the paired Lumber-minus-turn-packed recall and judge differences directly. Recall and judge scores are not combined into an arbitrary diagnostic score."),
        ("Interpretation", "The diagnostic strata support mechanism analysis. Only the seeded-random subset is a limited reality check; error frequencies over the full purposive sample are not prevalence estimates."),
        ("Reproduce", payload["command"]),
        ("Preset SHA-256", payload["hashes"]["preset"]),
        ("Retrieval summary SHA-256", payload["hashes"]["retrieval_summary"]),
        ("Answer summary SHA-256", payload["hashes"]["answer_summary"]),
        ("Evaluation SHA-256", payload["hashes"]["evaluation"]),
        ("Grounded sample SHA-256", payload["hashes"]["grounded_sample"]),
    )
    for label, explanation in rows:
        sheet.append([label, explanation])
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color="17365D")
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def write_summary(workbook, payload, annotation_end):
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet("Summary")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:T1")
    sheet["A1"] = "Qualitative analysis summary"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet["A2"] = (
        "Formula-driven summary. Complete the annotation fields on the Annotations "
        "sheet; the manual-analysis columns and charts update when opened in Excel."
    )
    sheet.merge_cells("A2:T2")
    sheet["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "Condition", "Chunker", "Budget", "N", "Precision", "Recall", "F1",
        "ROUGE-L", "BERTScore F1", "Judge", "Annotated", "Sufficient Yes",
        "Sufficient Partial", "Sufficient No", "Manual quality", "Unsupported",
        "High noise", "Auto faithfulness", "Auto sufficiency", "Auto unsupported",
    ]
    for column, value in enumerate(headers, 1):
        sheet.cell(4, column, value)
    row = 5
    for chunker in ANNOTATION_CHUNKERS:
        for budget in (ANNOTATION_BUDGET,):
            sheet.cell(row, 1, f"{chunker} / {budget}")
            sheet.cell(row, 2, chunker)
            sheet.cell(row, 3, budget)
            criteria = (
                f"'Annotations'!$F$2:$F${annotation_end},$B{row},"
                f"'Annotations'!$G$2:$G${annotation_end},$C{row}"
            )
            sheet.cell(row, 4, f"=COUNTIFS({criteria})")
            for column, source in enumerate(("I", "J", "K", "L", "M", "N"), 5):
                sheet.cell(
                    row,
                    column,
                    f'=IFERROR(AVERAGEIFS(\'Annotations\'!${source}$2:${source}${annotation_end},'
                    f"{criteria}),\"\")",
                )
            sheet.cell(
                row,
                11,
                f'=COUNTIFS({criteria},\'Annotations\'!$R$2:$R${annotation_end},"<>")',
            )
            for column, label in enumerate(("Yes", "Partial", "No"), 12):
                sheet.cell(
                    row,
                    column,
                    f'=IFERROR(COUNTIFS({criteria},\'Annotations\'!$R$2:$R${annotation_end},"{label}")/$K{row},"")',
                )
            sheet.cell(
                row,
                15,
                f'=IFERROR(AVERAGEIFS(\'Annotations\'!$T$2:$T${annotation_end},{criteria}),"")',
            )
            sheet.cell(
                row,
                16,
                f'=IFERROR(COUNTIFS({criteria},\'Annotations\'!$U$2:$U${annotation_end},"Yes")/$K{row},"")',
            )
            sheet.cell(
                row,
                17,
                f'=IFERROR(COUNTIFS({criteria},\'Annotations\'!$S$2:$S${annotation_end},"High")/$K{row},"")',
            )
            for column, source in ((18, "Y"), (19, "AB")):
                sheet.cell(
                    row,
                    column,
                    f'=IFERROR(AVERAGEIFS(\'Annotations\'!${source}$2:${source}${annotation_end},{criteria}),"")',
                )
            sheet.cell(
                row,
                20,
                f'=IFERROR(COUNTIFS({criteria},\'Annotations\'!$AA$2:$AA${annotation_end},"?*")/$D{row},"")',
            )
            row += 1

    sheet["A13"] = "Chunker comparison"
    chunker_headers = [
        "Chunker", "Recall", "Judge", "Manual quality", "Sufficient Yes",
        "Unsupported", "High noise",
    ]
    for column, value in enumerate(chunker_headers, 1):
        sheet.cell(14, column, value)
    for output_row, chunker in enumerate(ANNOTATION_CHUNKERS, 15):
        sheet.cell(output_row, 1, chunker)
        for column, source in ((2, "J"), (3, "N"), (4, "T")):
            sheet.cell(
                output_row,
                column,
                f'=IFERROR(AVERAGEIF(\'Annotations\'!$F$2:$F${annotation_end},$A{output_row},\'Annotations\'!${source}$2:${source}${annotation_end}),"")',
            )
        annotated = (
            f"COUNTIFS('Annotations'!$F$2:$F${annotation_end},$A{output_row},"
            f"'Annotations'!$R$2:$R${annotation_end},\"<>\")"
        )
        for column, source, label in (
            (5, "R", "Yes"), (6, "U", "Yes"), (7, "S", "High")
        ):
            sheet.cell(
                output_row,
                column,
                f'=IFERROR(COUNTIFS(\'Annotations\'!$F$2:$F${annotation_end},$A{output_row},\'Annotations\'!${source}$2:${source}${annotation_end},"{label}")/{annotated},"")',
            )

    sheet["A20"] = "Question sample"
    sheet["A21"], sheet["B21"] = "Query type", "Questions"
    for output_row, query_type in enumerate(QUERY_TYPES, 22):
        sheet.cell(output_row, 1, query_type)
        sheet.cell(
            output_row,
            2,
            f'=COUNTIF(\'Questions\'!$D$2:$D$31,A{output_row})',
        )

    sheet["A27"] = "Main error frequencies"
    sheet["A28"], sheet["B28"], sheet["C28"] = "Main error", "Count", "Share"
    errors = ("Retrieval", "Noise", "Generation", "Annotation", "None")
    for output_row, error in enumerate(errors, 29):
        sheet.cell(output_row, 1, error)
        sheet.cell(
            output_row,
            2,
            f'=COUNTIF(\'Annotations\'!$V$2:$V${annotation_end},A{output_row})',
        )
        sheet.cell(output_row, 3, f'=IFERROR(B{output_row}/SUM($B$29:$B$33),"")')

    sheet["A36"] = "Sampling and coding notes"
    notes = (
        f"30 questions: 10 per query type. Within each type: 2 largest Lumber "
        f"recall advantages, 2 largest turn-packed recall advantages, 2 Lumber "
        f"recall gains without a judge gain, and 4 seeded-random cases. Seed "
        f"{payload['seed']}; retriever {payload['retriever']}; 2,048 words.",
        "No composite diagnostic score is used. Sampling is based on paired Lumber-minus-turn-packed differences for the exact conditions shown in this workbook.",
        "Error frequencies across the full purposive sample describe reviewed cases; they must not be interpreted as population prevalence. The seeded-random subset is a limited reality check.",
        "Query type is assigned by a transparent lexical heuristic and is a sampling device, not a validated taxonomy.",
        "Synthesis: asks for a topic-level account spanning the meeting or a discussion segment.",
        "Participant-specific: names a meeting speaker and asks about that speaker's contribution or view.",
        "Reasoning/response: asks why/how, or about a decision, reaction, problem, solution, or outcome.",
        f"Retrieved-evidence answers come from {payload['answer_stage']}; the comparison answer comes from oracle-14b.",
        "Evidence sufficient: Yes = enough to answer fully; Partial = useful but incomplete; No = insufficient.",
        "Noise: Low = little irrelevant material; Med = noticeable but usable; High = distracting or misleading.",
        "Answer quality: 1 = invalid/incorrect; 2 = partially correct; 3 = correct.",
        "Unsupported content: the answer contains a claim not supported by retrieved evidence.",
        "Main error: choose the dominant cause only: Retrieval / Noise / Generation / Annotation / None.",
        f"Reproduce: {payload['command']}",
        f"Preset SHA-256: {payload['hashes']['preset']}",
        f"Retrieval summary SHA-256: {payload['hashes']['retrieval_summary']}",
        f"Answer summary SHA-256: {payload['hashes']['answer_summary']}",
        f"Evaluation SHA-256: {payload['hashes']['evaluation']}",
    )
    for output_row, note in enumerate(notes, 37):
        sheet.cell(output_row, 1, note)
        sheet.merge_cells(
            start_row=output_row, start_column=1, end_row=output_row, end_column=20
        )
        sheet.cell(output_row, 1).alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    header_fill = PatternFill("solid", fgColor="17365D")
    for header_row in (4, 14, 21, 28):
        for cell in sheet[header_row]:
            if cell.value is not None:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
    for title_cell in ("A13", "A20", "A27", "A36"):
        sheet[title_cell].font = Font(bold=True, color="17365D", size=12)
    sheet.freeze_panes = "A4"
    sheet.column_dimensions["A"].width = 28
    for column in range(2, 21):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    for output_row in range(5, 11):
        for column in range(5, 21):
            sheet.cell(output_row, column).number_format = "0.000"

    recall_chart = BarChart()
    recall_chart.title = "Retrieval recall by condition"
    recall_chart.add_data(
        Reference(sheet, min_col=6, min_row=4, max_row=6), titles_from_data=True
    )
    recall_chart.set_categories(Reference(sheet, min_col=1, min_row=5, max_row=6))
    recall_chart.height, recall_chart.width = 7, 12
    sheet.add_chart(recall_chart, "V4")

    quality_chart = BarChart()
    quality_chart.title = "Manual answer quality by condition"
    quality_chart.add_data(
        Reference(sheet, min_col=15, min_row=4, max_row=6), titles_from_data=True
    )
    quality_chart.set_categories(Reference(sheet, min_col=1, min_row=5, max_row=6))
    quality_chart.height, quality_chart.width = 7, 12
    sheet.add_chart(quality_chart, "V19")

    error_chart = PieChart()
    error_chart.title = "Main error frequencies"
    error_chart.add_data(
        Reference(sheet, min_col=2, min_row=28, max_row=33), titles_from_data=True
    )
    error_chart.set_categories(Reference(sheet, min_col=1, min_row=29, max_row=33))
    error_chart.height, error_chart.width = 7, 10
    sheet.add_chart(error_chart, "V34")


def existing_manual_annotations(path: Path) -> dict[tuple, tuple]:
    """Keep manually entered labels when refreshing the same review workbook."""
    if not path.exists():
        return {}
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook["Annotations"]
        saved = {}
        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=23, values_only=True):
            labels = tuple(row[17:23])
            if not any(value is not None and value != "" for value in labels):
                continue
            key = (row[1], row[2], row[5], row[6], row[7])
            if key in saved:
                raise ValueError(f"Duplicate annotated row in {path}: {key}")
            saved[key] = labels
        return saved
    finally:
        workbook.close()


def write_workbook(path: Path, payload: dict) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as error:
        raise RuntimeError("Install project dependencies to create the workbook") from error

    manual_annotations = existing_manual_annotations(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    write_instructions(workbook, payload)
    questions = workbook.create_sheet("Questions")
    question_headers = [
        "sample_id", "meeting_id", "question_index", "query_type",
        "selection_reason", "lumber_minus_turn_recall",
        "lumber_minus_turn_judge", "question", "reference_answer",
        "gold_turn_ranges", "gold_evidence", "gold_evidence_truncated",
        "oracle_14b_answer",
    ]
    question_rows = [
        [
            case["sample_id"], case["meeting_id"], case["question_index"],
            case["query_type"], case["selection_reason"], case["recall_delta"],
            case["judge_delta"], case["question"], case["reference_answer"],
            "; ".join(
                f"{start}-{end}" for start, end in case["gold_turn_ranges"]
            ),
            case["gold_evidence"], len(case["gold_evidence"]) > MAX_EXCEL_TEXT,
            case["oracle_answer"],
        ]
        for case in payload["selected"]
    ]
    add_table(
        questions,
        question_headers,
        question_rows,
        {"A": 10, "B": 14, "C": 12, "D": 20, "E": 38, "F": 18,
         "G": 18, "H": 48, "I": 58, "J": 18, "K": 100, "L": 18,
         "M": 85},
        "QuestionsTable",
    )
    for row in range(2, len(question_rows) + 2):
        questions.row_dimensions[row].height = 90
        questions.cell(row, 6).number_format = "0.000"
        questions.cell(row, 7).number_format = "0"

    annotations = workbook.create_sheet("Annotations")
    annotation_headers = [
        "sample_id", "meeting_id", "question_index", "query_type", "question",
        "chunker", "evidence_words", "retriever", "retrieval_precision",
        "retrieval_recall", "retrieval_f1", "rougeL", "bertscore_f1",
        "judge_score", "judge_reasoning", "retrieved_evidence", "generated_answer",
        "evidence_sufficient", "noise", "answer_quality_1_3",
        "unsupported_content", "main_error", "notes", "selection_reason",
        "auto_faithfulness_1_3", "auto_faithfulness_reason",
        "auto_unsupported_claims", "auto_sufficiency_1_3",
        "auto_sufficiency_reason", "auto_unavailable_reference_points",
    ]
    annotation_rows = []
    for case in payload["selected"]:
        for chunker in ANNOTATION_CHUNKERS:
            condition = (
                f"{chunker}__{payload['retriever']}__w{ANNOTATION_BUDGET}"
            )
            result = case["results"][condition]
            faithful, sufficient = result["faithfulness"], result["sufficiency"]
            precision, recall = result["precision"], result["recall"]
            f1 = (
                2 * precision * recall / (precision + recall)
                if precision + recall
                else 0.0
            )
            annotation_rows.append(
                [
                    case["sample_id"], case["meeting_id"],
                    case["question_index"], case["query_type"], case["question"],
                    chunker, ANNOTATION_BUDGET, payload["retriever"], precision,
                    recall, f1, result["rougeL"], result["bertscore_f1"],
                    result["judge"], result["judge_reason"],
                    result["retrieved_evidence"], result["answer"], "", "", "",
                    "", "", "", case["selection_reason"],
                    faithful["score"], faithful["reason"],
                    "\n".join(faithful["unsupported_claims"]),
                    sufficient["score"], sufficient["reason"],
                    "\n".join(sufficient["unavailable_reference_points"]),
                ]
            )
            key = (
                case["meeting_id"], case["question_index"], chunker,
                ANNOTATION_BUDGET, payload["retriever"],
            )
            if key in manual_annotations:
                annotation_rows[-1][17:23] = manual_annotations.pop(key)
    if manual_annotations:
        raise ValueError(
            f"Existing manual annotations would be lost: {list(manual_annotations)}"
        )
    add_table(
        annotations,
        annotation_headers,
        annotation_rows,
        {"A": 10, "B": 14, "C": 12, "D": 20, "E": 44, "F": 16,
         "G": 14, "H": 12, "I": 15, "J": 14, "K": 12, "L": 12,
         "M": 15, "N": 12, "O": 42, "P": 85, "Q": 65, "R": 18,
         "S": 12, "T": 18, "U": 20, "V": 16, "W": 42, "X": 38,
         "Y": 16, "Z": 45, "AA": 60, "AB": 16, "AC": 45, "AD": 60},
        "AnnotationsTable",
    )
    annotation_end = len(annotation_rows) + 1
    for row in range(2, annotation_end + 1):
        annotations.row_dimensions[row].height = 95
    for column in ("I", "J", "K", "L", "M"):
        for cell in annotations[column][1:]:
            cell.number_format = "0.000"
    validations = {
        "R": '"Yes,Partial,No"',
        "S": '"Low,Med,High"',
        "U": '"Yes,No"',
        "V": '"Retrieval,Noise,Generation,Annotation,None"',
    }
    for column, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula)
        annotations.add_data_validation(validation)
        validation.add(f"{column}2:{column}{annotation_end}")
    quality = DataValidation(
        type="whole", operator="between", formula1="1", formula2="3"
    )
    annotations.add_data_validation(quality)
    quality.add(f"T2:T{annotation_end}")
    annotations.conditional_formatting.add(
        f"T2:T{annotation_end}",
        CellIsRule(
            operator="equal",
            formula=["1"],
            fill=PatternFill("solid", fgColor="F4CCCC"),
        ),
    )
    annotations.conditional_formatting.add(
        f"T2:T{annotation_end}",
        CellIsRule(
            operator="equal",
            formula=["3"],
            fill=PatternFill("solid", fgColor="D9EAD3"),
        ),
    )

    write_summary(workbook, payload, annotation_end)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--preset", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--retriever", default="dense")
    parser.add_argument("--questions-per-type", type=int, default=10)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--answer-stage", default="retrieval-14b")
    args = parser.parse_args()

    run = load_run_config(args.preset)
    if args.retriever not in run.retrieval.retrievers:
        raise ValueError(f"Retriever is absent from preset: {args.retriever}")
    if ANNOTATION_BUDGET not in run.retrieval.evidence_budgets:
        raise ValueError(f"Preset must include budget {ANNOTATION_BUDGET}")
    conditions = [
        f"{chunker}__{args.retriever}__w{ANNOTATION_BUDGET}"
        for chunker in ANNOTATION_CHUNKERS
    ]
    cases = load_cases(run, args.answer_stage, args.retriever)
    add_sampling_deltas(cases, args.retriever)
    selected = select_cases(cases, args.questions_per_type, args.seed)
    hydrate_evidence(run, selected, conditions)
    grounded_sample = run.output_root / "grounded-judge" / "sample.json"
    add_evidence_aware_judgments(selected, grounded_sample, args.retriever)

    command = (
        f"python src/tools/export_qualitative_workbook.py --preset "
        f"{args.preset.as_posix()} --output {args.output.as_posix()} --retriever "
        f"{args.retriever} --questions-per-type {args.questions_per_type} "
        f"--seed {args.seed}"
    )
    payload = {
        "selected": selected,
        "retriever": args.retriever,
        "answer_stage": args.answer_stage,
        "seed": args.seed,
        "command": command,
        "hashes": {
            "preset": sha256(args.preset),
            "retrieval_summary": sha256(run.retrieval_dir / "summary.json"),
            "answer_summary": sha256(
                run.answers_dir / args.answer_stage / "summary.json"
            ),
            "evaluation": sha256(
                run.evaluation_dir / f"{args.answer_stage}.json"
            ),
            "grounded_sample": sha256(grounded_sample),
        },
    }
    write_workbook(args.output, payload)
    print(f"Qualitative review workbook: {args.output}")


if __name__ == "__main__":
    main()
